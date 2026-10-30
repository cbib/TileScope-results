"""Check the submission data round trip, not the incidental directory layout."""
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook

from analyses.figures import common


def test_panel_data_round_trip_keeps_ids_missing_values_and_manual_layout(tmp_path, monkeypatch):
    monkeypatch.setattr(common, 'ROOT', tmp_path)
    out=common.figure_dir('Figure1')
    assembly=out/'assembly';assembly.mkdir(parents=True)
    manual=assembly/'author_layout.svg';manual.write_text('author-owned layout')
    fig,ax=plt.subplots();ax.plot([0,1],[.2,.4])
    common.save_panel(fig,out/'panels','A')
    values=pd.DataFrame({'roi_id':['001','=not_a_formula'],'r':[.2,np.nan],
                         'status':['finite','unavailable']})
    geometry=pd.DataFrame({'wkt':['x'*32768]})
    common.write_sources('Figure1',{'values':values,'geometry':geometry},
                         {'A':['values','geometry']},notes='Values are Pearson coefficients.')
    workbook=load_workbook(out/'source_data/Figure1_Source_Data.xlsx',data_only=False)
    assert workbook['D01']['A2'].value=='001'
    assert workbook['D01']['A3'].value=='=not_a_formula'
    assert workbook['D01']['A3'].data_type=='s'
    assert workbook['D01']['B3'].value is None
    assert workbook['D01']['C3'].value=='unavailable'
    assert 'D02' not in workbook.sheetnames
    workbook.close()
    pd.testing.assert_frame_equal(pd.read_csv(out/'source_data/values.csv',dtype={'roi_id':str}),values)
    assert pd.read_csv(out/'source_data/geometry.csv').wkt.iloc[0]=='x'*32768
    index=pd.read_csv(out/'source_data/tables.csv',keep_default_na=False)
    assert index.panels.tolist()==['A','A']
    assert 'without truncation' in index.note.iloc[1]
    assert manual.read_text()=='author-owned layout'
