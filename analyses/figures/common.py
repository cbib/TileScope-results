"""Small shared export utilities for manuscript panels; never runs TileScope."""
from __future__ import annotations

import hashlib
import html
import json
import os
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from tilescope_results.paths import REPOSITORY, reference

ROOT = Path(os.environ.get('TILESCOPE_FIGURE_DIR', REPOSITORY / 'analyses/figures')).expanduser().resolve()
EXPRESSION_COLOR_STOPS = ((0., '#80B5D6'), (.25, '#C5DBE9'), (.5, '#F6F3EE'),
                          (.75, '#EC947F'), (1., '#B91F2D'))
BANDS = {'anchor':'#343c48', 'near':'#d47d2e', 'buffer':'#e8cf8b',
         'far':'#6a9dbd', 'unreachable':'#bebebe'}


def figure_dir(name):
    if name in {'Mouse', 'Benchmark'}:
        return ROOT / 'additional' / name
    return ROOT / ('supplementary' if name.startswith('FigureS') else 'main') / name


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def expression_colormap():
    from matplotlib.colors import LinearSegmentedColormap
    return LinearSegmentedColormap.from_list('figure1_light_blue_red', EXPRESSION_COLOR_STOPS)


def spot_label_outlines(frame, spacing, rois):
    """Display-only Voronoi footprints; never substitutes for analysis support."""
    from scipy.spatial import Voronoi
    from shapely.geometry import Point, Polygon
    from shapely.ops import unary_union
    xy=frame[['x','y']].to_numpy()
    lo,hi=xy.min(axis=0),xy.max(axis=0)
    centre,span=(lo+hi)/2,float(max(hi-lo))
    guards=centre+4*span*np.array([[-1,-1],[-1,1],[1,-1],[1,1]])
    vor=Voronoi(np.vstack([xy,guards]))
    result={}
    for roi in rois:
        cells=[]
        for i in np.flatnonzero(frame.roi_id.eq(roi)):
            vertices=vor.regions[vor.point_region[i]]
            if not vertices or -1 in vertices: raise ValueError('Unbounded observed footprint')
            cell=Polygon(vor.vertices[vertices]).intersection(Point(xy[i]).buffer(.65*spacing,quad_segs=8))
            if not cell.is_valid or not cell.covers(Point(xy[i])): raise ValueError('Footprint misses source spot')
            cells.append(cell)
        result[roi]=unary_union(cells)
    return result


def save_panel(fig, directory, stem):
    """SVG/PDF are editable artwork; PNG is a white-background preview."""
    directory.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({'svg.fonttype':'none', 'pdf.fonttype':42})
    paths = []
    for ext in ('svg', 'pdf', 'png'):
        path = directory / f'{stem}.{ext}'
        meta = {'Date':None} if ext == 'svg' else ({'CreationDate':None, 'ModDate':None} if ext == 'pdf' else None)
        fig.savefig(path, dpi=300, transparent=ext != 'png', facecolor='white' if ext == 'png' else 'none', metadata=meta)
        paths.append(path)
    plt.close(fig)
    return paths


def write_sources(name, tables, panel_tables, *, inputs=(), notes='', output=None):
    """Write the same frames to CSV and Excel, with explicit panel mapping.

    Geometry strings beyond Excel's cell limit remain lossless in CSV; their
    sheets are explicitly omitted, never truncated. IDs and other strings are
    written as strings, not formulas or automatically converted hyperlinks.
    """
    out = Path(output) if output is not None else figure_dir(name)
    source = out / 'source_data'
    source.mkdir(parents=True, exist_ok=True)
    records = []
    for key, frame in tables.items():
        if not isinstance(frame, pd.DataFrame):
            raise TypeError(f'{key} must be a DataFrame')
        file = source / f'{key}.csv'
        frame.to_csv(file, index=False)
        panels = [p for p, keys in panel_tables.items() if key in keys]
        long_cells = any(frame[c].astype(str).str.len().gt(32767).any()
                         for c in frame.select_dtypes(include=['object', 'string']).columns)
        records.append({'table':key, 'panels':'; '.join(panels), 'rows':len(frame),
                        'csv':file.name, 'excel_sheet':f'D{len(records)+1:02d}' if not long_cells else '',
                        'note':'Geometry exceeds Excel cell limit; use CSV without truncation.' if long_cells else ''})
    index = pd.DataFrame(records)
    workbook = source / f'{name}_Source_Data.xlsx'
    with pd.ExcelWriter(workbook, engine='xlsxwriter', engine_kwargs={'options':{
            'strings_to_formulas':False, 'strings_to_urls':False}}) as writer:
        pd.DataFrame({'Figure':[name], 'Interpretation':[notes], 'Missing values':[
            'Blank numeric cells mean unavailable/uncomputed, never measured zero; consult status columns.'],
            'Data unit':['IDs retain sample, ROI, component, interface, replicate or spatial-unit grain.']}).to_excel(
                writer, sheet_name='README', index=False)
        index.to_excel(writer, sheet_name='Panel_index', index=False)
        for row in records:
            if row['excel_sheet']:
                tables[row['table']].to_excel(writer, sheet_name=row['excel_sheet'], index=False)
        for sheet in writer.sheets.values():
            sheet.freeze_panes(1, 0)
            sheet.set_column(0, 0, 24)
    # Verify that the exported workbook opens with a separate Excel reader.
    from openpyxl import load_workbook
    check = load_workbook(workbook, read_only=True, data_only=False)
    for row in records:
        if row['excel_sheet'] and check[row['excel_sheet']].max_row != row['rows'] + 1:
            raise ValueError(f'Workbook row loss: {name}/{row["table"]}')
    check.close()
    index.to_csv(source / 'tables.csv', index=False)
    (source / 'README.md').write_text(
        f'# {name} source data\n\n{notes}\n\n'
        'Use `tables.csv` (or the workbook Panel_index sheet) to find each panel’s tables. '
        'CSV and Excel contain the same values; geometry exceeding the Excel cell limit is '
        'provided in full in CSV and explicitly identified in the index. Blank numeric '
        'cells are unavailable/uncomputed, not zero. Keep status columns when filtering.\n\n'
        'Sample/ROI/component/interface identifiers define the observational units; spots, '
        'hexagons and tessellations are not independent biological replicates. Pearson and '
        'Moran coefficients, enrichment ratios and distance profiles retain their original '
        'scales. Expression maps use the recorded input scores and coordinates. '
        'Units, display settings and denominators are documented above and in the source tables. '
        'Figure captions are maintained in the manuscript.\n'
    )
    panel_rows=[]
    for stem, keys in panel_tables.items():
        missing=set(keys)-set(tables)
        if missing: raise ValueError(f'{stem}: missing source tables {missing}')
        for ext in ('svg','pdf','png'):
            if not (out/'panels'/f'{stem}.{ext}').is_file():
                raise ValueError(f'Missing panel {name}/{stem}.{ext}')
        panel_rows.append(dict(figure=name,panel=stem,svg=f'panels/{stem}.svg',pdf=f'panels/{stem}.pdf',
                               source_data=';'.join(f'source_data/{key}.csv' for key in keys)))
    pd.DataFrame(panel_rows).to_csv(out/'panel_index.csv',index=False)
    (out/'provenance.json').write_text(json.dumps({
        'analysis_rerun':False, 'inputs':{reference(Path(p)):digest(p) for p in sorted(set(inputs))},
        'generators':{reference(p):digest(p) for p in sorted((REPOSITORY/'analyses/figures/scripts').rglob('*.py'))},
        'exporter_sha256':digest(Path(__file__)),
        'outputs':{str(p.relative_to(out)):digest(p) for folder in ('panels','source_data')
                   for p in sorted((out/folder).glob('*')) if p.is_file()}},indent=2)+'\n')
    return panel_rows


def write_index():
    records=[]
    for path in sorted(ROOT.glob('*/*/panel_index.csv')):
        rows=pd.read_csv(path,keep_default_na=False)
        prefix=path.parent.relative_to(ROOT)
        for row in rows.to_dict('records'):
            for key in ('svg','pdf'):
                row[key]=str(prefix/row[key])
            row['source_data']=';'.join(str(prefix/p) for p in row['source_data'].split(';'))
            records.append(row)
    pd.DataFrame(records).to_csv(ROOT/'index.csv',index=False)
    cards=[]
    for row in records:
        svg=html.escape(row['svg']); pdf=html.escape(row['pdf'])
        source=html.escape(row['source_data'].split(';')[0])
        cards.append(f'<article><h3>{html.escape(row["figure"]+" / "+row["panel"])}</h3>'
                     f'<a href="{svg}">SVG</a> · <a href="{pdf}">PDF</a> · <a href="{source}">Data</a>'
                     f'<img loading="lazy" src="{svg}"></article>')
    (ROOT/'index.html').write_text('<!doctype html><meta charset="utf-8"><title>TileScope figure panels</title>'
        '<style>body{font:16px system-ui;margin:2em}main{columns:2}article{break-inside:avoid;margin-bottom:2em}'
        'img{display:block;width:100%;max-height:550px;object-fit:contain}</style><h1>Individual figure panels</h1>'
        '<p>Editable artwork and linked source data. Final layouts are assembled separately.</p><main>'+
        ''.join(cards)+'</main>')
    print(f'Indexed {len(records)} individual panels in {ROOT}',flush=True)


def rebuild_bands(frame, config, anchor, responder):
    """Reconstruct the existing core bands for display; not a new estimator."""
    from tilescope.spatial.source_spot_table import build_source_spot_table_from_frame
    from tilescope.spatial.source_spot_adjacency import build_source_spot_adjacency, UNREACHABLE_HOP
    from tilescope.analyzers.parent_roi_interface_response import NEAR_HOP_MAX, BUFFER_HOP, FAR_HOP_MIN
    if config['run_config']['technology'] != 'visium': raise ValueError('Expected Visium support')
    policy=config['resolved_policy']['source_adjacency']
    if policy['method'] != 'fixed_radius': raise ValueError('Unexpected source graph')
    graph=build_source_spot_adjacency(build_source_spot_table_from_frame(frame),technology='visium',radius_factor=policy['radius_factor'])
    _,seeds=graph.interface_endpoints(anchor,responder)
    hops=graph.hop_distance_within_roi(responder,seeds)
    inside=frame.roi_id.eq(responder).to_numpy()
    bands=np.full(len(frame),'outside',dtype=object)
    bands[frame.roi_id.eq(anchor)]='anchor'
    bands[inside&(hops==UNREACHABLE_HOP)]='unreachable'
    bands[inside&(hops>=0)&(hops<=NEAR_HOP_MAX)]='near'
    bands[inside&(hops==BUFFER_HOP)]='buffer'
    bands[inside&(hops>=FAR_HOP_MIN)]='far'
    if (bands[inside]=='outside').any(): raise ValueError('Unclassified responder hop')
    return bands,hops,graph.cross_roi_edges(anchor,responder)
